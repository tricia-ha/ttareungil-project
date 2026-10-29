"""Official operator timetable adapters for additional metropolitan rail lines.

Train-matrix sources preserve train identities and arrival/departure seconds.
Station-only sources require explicit, reported train matching; they are not
represented as operator-issued train-level schedules.
"""

from pathlib import Path
from collections import Counter, defaultdict
from datetime import time, timedelta
import csv
import json
import re
import hashlib
import math
from functools import lru_cache

LINES = (
    "경의중앙선",
    "수인분당선",
    "경춘선",
    "공항철도",
    "신분당선",
    "우이신설선",
    "신림선",
)
FIELDS = (
    "LINE_NUM",
    "STATION_CD",
    "STATION_NM",
    "TRAIN_NO",
    "WEEK_TAG",
    "INOUT_TAG",
    "BRANCH_LINE",
    "ORIGINSTATION",
    "DESTSTATION",
    "ARRIVETIME",
    "LEFTTIME",
    "VALID_FROM",
    "VALID_UNTIL",
    "TIMETABLE_METHOD",
)


def name(value):
    text = (
        re.sub(r"\([^)]*\)", "", str(value or ""))
        .replace(" ", "")
        .replace("·", "")
        .replace("ㆍ", "")
    )
    if text.endswith("역"):
        text = text[:-1]
    # Abbreviations in the operator's train matrix, checked against station order.
    return {
        "서울": "서울",
        "4.19민주묘지": "419민주묘지",
        "디엠씨": "디지털미디어시티",
        "디엠시": "디지털미디어시티",
        "평내호": "평내호평",
        "효창공": "효창공원앞",
        "홍대입": "홍대입구",
        "항공대": "한국항공대",
        "1양정": "양정",
        "1양원": "양원",
        "대모산": "대모산입구",
        "강남구": "강남구청",
        "로데오": "압구정로데오",
        "신수원": "수원",
        "수원시": "수원시청",
        "매탄권": "매탄권선",
        "신인천": "인천",
        "남동인": "남동인더스파크",
        "인천논": "인천논현",
        "소래포": "소래포구",
        "신길온": "능길",
        "신길온천": "능길",
        "인천공항": "인천공항1터미널",
    }.get(text, text)


def clock(value):
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, timedelta):
        value = int(value.total_seconds())
        return f"{value // 3600:02d}:{value // 60 % 60:02d}:{value % 60:02d}"
    text = str(value or "").strip()
    return text if re.fullmatch(r"\d{1,2}:\d{2}:\d{2}", text) else ""


def hms(seconds):
    return f"{seconds // 3600:02d}:{seconds // 60 % 60:02d}:{seconds % 60:02d}"


@lru_cache(maxsize=4)
def station_catalog(folder):
    from openpyxl import load_workbook

    workbook = load_workbook(
        Path(folder) / "files" / "national_rail_stations.xlsx",
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    aliases = {
        "분당선": "수인분당선",
        "수인선": "수인분당선",
        "인천국제공항선": "공항철도",
        "수도권경량도시철도신림선": "신림선",
    }
    options = defaultdict(list)
    korail_points = defaultdict(list)
    for values in rows:
        r = dict(zip(headers, values))
        line = str(r["노선명"]).replace(" ", "")
        line = aliases.get(line, line)
        try:
            lat, lon = float(r["역위도"]), float(r["역경도"])
        except (ValueError, TypeError):
            continue
        if not (33 < lat < 39 and 124 < lon < 132):
            continue
        point = {"name": str(r["역사명"]), "latitude": lat, "longitude": lon}
        if line in LINES:
            options[(line, name(r["역사명"]))].append(point)
        if 36.8 < lat < 38.2 and 126.4 < lon < 128.2:
            options[("*", name(r["역사명"]))].append(point)
        if str(r.get("운영기관명", "")) == "한국철도공사":
            korail_points[name(r["역사명"])].append(point)
    workbook.close()
    for st, points in korail_points.items():
        options[("*", st)] = points
    return {
        key: values[0]
        for key, values in options.items()
        if all(
            abs(v["latitude"] - values[0]["latitude"]) < 0.003
            and abs(v["longitude"] - values[0]["longitude"]) < 0.003
            for v in values
        )
    }


def matrix_rows(path, line_for_sheet, day_for_sheet, valid_from, valid_until=""):
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in workbook:
            line = line_for_sheet(sheet.title)
            if not line:
                continue
            days = day_for_sheet(sheet.title)
            data = list(sheet.iter_rows(values_only=True))
            headers = [
                i for i, r in enumerate(data) if r and str(r[0]).strip() == "열차번호"
            ]
            for block, header in enumerate(headers):
                end = headers[block + 1] - 1 if block + 1 < len(headers) else len(data)
                origins = next(
                    (
                        r
                        for r in data[max(0, header - 3) : header + 5]
                        if str(r[0]).strip() in ("시발역", "시작역")
                    ),
                    None,
                )
                destinations = next(
                    (
                        r
                        for r in data[max(0, header - 3) : header + 5]
                        if str(r[0]).strip() in ("종착역", "종착")
                    ),
                    None,
                )
                for col, number in enumerate(data[header][1:], 1):
                    train = str(number or "").strip()
                    if not re.fullmatch(r"[A-Z]?\d+", train):
                        continue
                    # AREX A1xxx are reserved-seat express trains, not local metro service.
                    if line == "공항철도" and train.startswith("A1"):
                        continue
                    origin = (
                        name(origins[col]) if origins and col < len(origins) else ""
                    )
                    terminal = (
                        name(destinations[col])
                        if destinations and col < len(destinations)
                        else ""
                    )
                    events = []
                    for i in range(header + 1, end - 1):
                        row = data[i]
                        following = data[i + 1]
                        st = name(row[0])
                        if (
                            not st
                            or str(following[0] or "").strip()
                            or col >= len(following)
                        ):
                            continue
                        arrival = clock(row[col]) if col < len(row) else ""
                        departure = clock(following[col])
                        if not arrival and not departure:
                            continue
                        # A departure-only intermediate entry can be a passing time.
                        if not arrival and origin and st != origin:
                            continue
                        if not departure and terminal and st != terminal:
                            continue
                        arrival = arrival or departure
                        departure = departure or arrival
                        events.append((st, arrival, departure))
                    if len(events) < 2:
                        continue
                    for tag in days:
                        for st, arr, dep in events:
                            yield dict(
                                zip(
                                    FIELDS,
                                    [
                                        line,
                                        line + ":" + st,
                                        st,
                                        train,
                                        tag,
                                        sheet.title,
                                        sheet.title + ":" + str(block),
                                        origin,
                                        terminal,
                                        arr,
                                        dep,
                                        valid_from,
                                        valid_until,
                                        "official_train_matrix",
                                    ],
                                )
                            )
    finally:
        workbook.close()


def station_tables(folder, line):
    from bs4 import BeautifulSoup

    counts = {"신분당선": 16, "우이신설선": 13, "신림선": 11}
    prefix = {"신분당선": "shinbundang", "우이신설선": "ui", "신림선": "sillim"}[line]
    catalog = station_catalog(Path(folder).parent)
    all_stations = []
    for index in range(1, counts[line] + 1):
        path = Path(folder) / f"{prefix}_{index:02d}.html"
        soup = BeautifulSoup(path.read_bytes(), "html.parser")
        if line == "우이신설선":
            selected = soup.select_one('select[name="selST"] option[selected]')
            st = name(selected.get_text()) if selected else ""
            tables = [t for t in soup.find_all("table") if "시간(분)" in t.get_text()]
        elif line == "신림선":
            st = name(soup.select_one("#tab_moType1 li.on").get_text())
            tables = [t for t in soup.find_all("table") if "시, 분" in t.get_text()]
        else:
            # Operator order is stable and the route page carries station names.
            route = BeautifulSoup(
                (Path(folder) / "discovery" / "shinbundang.html").read_bytes(),
                "html.parser",
            )
            st = name(route.select_one(f"#SB{index:03d} strong").get_text())
            tables = soup.select("table.tb_st01")
        if (line, st) not in catalog:
            raise ValueError(f"추가 노선 좌표 미매칭: {line} {st}")
        events = defaultdict(list)
        for ti, table in enumerate(tables):
            for tr in table.select("tbody tr"):
                cells = tr.find_all(["td", "th"], recursive=False)
                if not cells:
                    continue
                if line == "신림선":
                    if len(cells) != 2:
                        continue
                    hour_text = cells[0].get_text()
                    sides = [(ti % 2, cells[1])]
                    day = ti // 2
                else:
                    if len(cells) != 3:
                        continue
                    hour_text = cells[1].get_text()
                    sides = [(0, cells[0]), (1, cells[2])]
                    day = ti
                match = re.search(r"\d+", hour_text)
                if not match:
                    continue
                hour = int(match[0])
                hour = hour + 24 if hour < 3 else hour
                for direction, cell in sides:
                    for minute, terminal in re.findall(
                        r"(?<!\d)([0-5]\d)(?:\(([^)]+)\))?",
                        cell.get_text(" ", strip=True),
                    ):
                        events[(day, direction)].append(
                            (hour * 3600 + int(minute) * 60, name(terminal))
                        )
        all_stations.append(
            (st, {key: sorted(set(value)) for key, value in events.items()})
        )
    return all_stations


def inferred_station_rows(folder, line):
    """Match consecutive minute-resolution station departures conservatively.

    The operator does not publish train IDs in these tables. The matching is
    marked inferred, never passed off as an exact train-level timetable.
    """
    stations = station_tables(folder, line)
    catalog = station_catalog(Path(folder).parent)
    from multimodal_data import distance

    for day, tags in ((0, ("1",)), (1, ("2", "3"))):
        for direction in (0, 1):
            # Sillim table 0 is down, the other operators table 0 is up.
            forward = (direction == 0) if line == "신림선" else (direction == 1)
            ordered = stations if forward else stations[::-1]
            tracks = []
            last_events = []
            for si, (station, events) in enumerate(ordered):
                departures = events.get((day, direction), [])
                if not departures:
                    # Terminal arrival times are not supplied by departure tables.
                    # Keep the preceding segment unavailable instead of inventing times.
                    continue
                used = set()
                current = []
                if si and last_events:
                    straight = distance(
                        catalog[(line, ordered[si - 1][0])], catalog[(line, station)]
                    )
                    # Reject pairings faster than 126 km/h even before bends/dwell.
                    minimum = max(60, math.ceil(straight / 35 / 60) * 60)
                    maximum = max(600, minimum + 300)
                    # Early low-frequency service identifies the local travel-time
                    # phase more reliably than peak-period repeated intervals.
                    early = [(t, end) for t, end, _ in last_events if t < 7 * 3600]
                    deltas = [
                        b - a
                        for a, _ in early
                        for b, _ in departures
                        if minimum <= b - a <= maximum
                    ]
                    offset = Counter(deltas).most_common(1)[0][0] if deltas else 120
                    last_index = -1
                    for t, end, track in last_events:
                        if end and name(ordered[si - 1][0]) == end:
                            continue
                        choices = [
                            (abs((dt - t) - offset), j, dt, dend)
                            for j, (dt, dend) in enumerate(departures)
                            if j > last_index
                            and j not in used
                            and minimum <= dt - t <= maximum
                            and abs((dt - t) - offset) <= 60
                            and (not end or not dend or end == dend)
                        ]
                        if not choices:
                            continue
                        _, j, dt, dend = min(choices)
                        used.add(j)
                        last_index = j
                        tracks[track].append((station, dt, dend))
                        current.append((dt, dend, track))
                for j, (dt, end) in enumerate(departures):
                    if j in used:
                        continue
                    track = len(tracks)
                    tracks.append([(station, dt, end)])
                    current.append((dt, end, track))
                last_events = sorted(current)
            for track, events in enumerate(tracks):
                if len(events) < 2:
                    continue
                terminals = {end for _, _, end in events if end}
                if len(terminals) > 1:
                    continue
                terminal = next(iter(terminals)) if terminals else ordered[-1][0]
                for tag in tags:
                    for st, seconds, end in events:
                        yield dict(
                            zip(
                                FIELDS,
                                [
                                    line,
                                    line + ":" + st,
                                    st,
                                    f"inferred-{direction}-{track}",
                                    tag,
                                    str(direction),
                                    "station_matching",
                                    events[0][0],
                                    terminal,
                                    hms(seconds),
                                    hms(seconds),
                                    "2026-09-28"
                                    if line == "신분당선"
                                    else "2026-10-17",
                                    "",
                                    "inferred_minute_station_matching",
                                ],
                            )
                        )


def prepare_extended_rail(folder, progress=print):
    folder = Path(folder)
    source = folder / "extended_rail"
    if not source.is_dir():
        return None
    required = [source / f"korail_{n}.xlsx" for n in (25745, 25746, 25702)] + [
        source / "arex_20251229.xlsx",
        source / "arex_20261023.xlsx",
    ]
    required += [
        source / f"{prefix}_{i:02d}.html"
        for prefix, count in (("shinbundang", 16), ("ui", 13), ("sillim", 11))
        for i in range(1, count + 1)
    ]
    required += [
        source / "discovery" / "shinbundang.html",
        folder / "files" / "national_rail_stations.xlsx",
    ]
    missing = [str(p) for p in required if not p.is_file()]
    if missing:
        raise FileNotFoundError(
            "수도권 추가 시간표 원자료 미확보: " + ", ".join(missing[:3])
        )
    fingerprint = hashlib.sha256(b"extended-rail-v1")
    fingerprint.update(Path(__file__).read_bytes())
    for p in required:
        fingerprint.update(p.read_bytes())
    signature = fingerprint.hexdigest()
    target = source / "normalized.csv"
    manifest = source / "manifest.json"
    if target.is_file() and manifest.is_file():
        cached = json.loads(manifest.read_text(encoding="utf-8"))
        if cached.get("signature") == signature:
            return cached
    station_catalog.cache_clear()
    catalog = station_catalog(folder)

    def korail_line(title):
        if title.startswith(("경의중앙", "경의선")):
            return "경의중앙선"
        if title.startswith(("경춘선", "망우직결")):
            return "경춘선"
        return None

    all_rows = []
    for idx, tags in ((25745, ("1",)), (25746, ("2", "3"))):
        if progress:
            progress(f"추가 시간표 변환: 코레일 {idx}", flush=True)
        all_rows.extend(
            matrix_rows(
                source / f"korail_{idx}.xlsx", korail_line, lambda s: tags, "2026-09-01"
            )
        )
    all_rows.extend(
        matrix_rows(
            source / "korail_25702.xlsx",
            lambda s: "수인분당선",
            lambda s: ("1",) if "평일" in s else ("2", "3"),
            "2026-08-22",
        )
    )
    for filename, start, end in (
        ("arex_20251229.xlsx", "2025-12-29", "2026-10-22"),
        ("arex_20261023.xlsx", "2026-10-23", ""),
    ):
        if progress:
            progress("추가 시간표 변환: " + filename, flush=True)
        all_rows.extend(
            matrix_rows(
                source / filename,
                lambda s: "공항철도",
                lambda s: ("1",) if "평일" in s else ("2", "3"),
                start,
                end,
            )
        )
    for line in ("신분당선", "우이신설선", "신림선"):
        if progress:
            progress("추가 시간표 연결 검증: " + line, flush=True)
        all_rows.extend(inferred_station_rows(source, line))
    stops = {}
    unmatched = Counter()
    for row in all_rows:
        key = (row["LINE_NUM"], name(row["STATION_NM"]))
        if key not in catalog and ("*", key[1]) in catalog:
            catalog[key] = catalog[("*", key[1])]
        if key not in catalog:
            unmatched[key] += 1
            continue
        stops[row["STATION_CD"]] = catalog[key]
    # An unresolved station invalidates its trip, not just a row. This prevents
    # skipping an unknown stop and silently turning it into an express segment.
    invalid = {
        tuple(
            r[k]
            for k in ("LINE_NUM", "TRAIN_NO", "WEEK_TAG", "BRANCH_LINE", "VALID_FROM")
        )
        for r in all_rows
        if r["STATION_CD"] not in stops
    }
    all_rows = [
        r
        for r in all_rows
        if tuple(
            r[k]
            for k in ("LINE_NUM", "TRAIN_NO", "WEEK_TAG", "BRANCH_LINE", "VALID_FROM")
        )
        not in invalid
    ]
    report = {
        "signature": signature,
        "rows_file": "extended_rail/normalized.csv",
        "rows_by_line": dict(Counter(r["LINE_NUM"] for r in all_rows)),
        "rows_by_method": dict(Counter(r["TIMETABLE_METHOD"] for r in all_rows)),
        "stops": stops,
        "unmatched_stations": [
            {"line": k[0], "name": k[1], "rows": v} for k, v in unmatched.items()
        ],
        "invalid_trips_missing_coordinates": len(invalid),
        "limitations": [
            "신분당선·우이신설선·신림선은 분 단위 역별 출발시각으로 열차 연결을 추정함",
            "역별 출발시각만 있는 노선은 종착역 도착시각이 없어 마지막 구간을 생성하지 않음",
        ],
        "sources": {
            "korail": "https://www.korail.com/com/userBoard.do?mode=list&schBcid=ticketTable",
            "arex": "https://www.airportrailroad.com/train/normal/info/010/0",
            "shinbundang": "https://www.shinbundang.co.kr/dxline/dxline3_1.jsp",
            "ui": "https://www.ui-line.com/html/intro/intro00/intro_00_01.php",
            "sillim": "https://www.sillimlrt.com/kr/html/sub01/01010201.html",
        },
    }
    part = target.with_suffix(".csv.part")
    with part.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(all_rows)
    part.replace(target)
    manifest.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    if progress:
        progress("추가 수도권 시간표 준비: " + str(report["rows_by_line"]), flush=True)
    return report


def read_extended_rows(folder, tags=None):
    path = Path(folder) / "extended_rail" / "normalized.csv"
    if path.is_file():
        with path.open(encoding="utf-8-sig", newline="") as stream:
            for row in csv.DictReader(stream):
                if tags is None or row["WEEK_TAG"] in tags:
                    yield row
